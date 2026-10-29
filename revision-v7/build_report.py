"""Build the standalone school submission with every original project photograph."""
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
COPY = json.loads((ROOT / 'report-copy.json').read_text(encoding='utf-8'))
TARGET = ROOT.parent / 'Rocket-Complete-Report-v7.docx'
PHOTO_ORDER = []
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
section.footer_distance = Inches(.35)
section.different_first_page_header_footer = True
normal = document.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
normal.element.get_or_add_rPr().append(OxmlElement('w:lang'))
normal.element.rPr.find(qn('w:lang')).set(qn('w:val'), 'es-CO')
document.core_properties.author = COPY['author']
document.core_properties.title = COPY['title'] + ' — ' + COPY['subtitle']
document.core_properties.subject = COPY['subject']
document.core_properties.language = 'es-CO'
settings = document.settings.element
compression = OxmlElement('w:doNotAutoCompressPictures')
settings.append(compression)


def paragraph(text, *, size=11, bold=False, center=False, before=0, after=7, keep=False):
    item = document.add_paragraph()
    item.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    item.paragraph_format.space_before = Pt(before)
    item.paragraph_format.space_after = Pt(after)
    item.paragraph_format.keep_with_next = keep
    run = item.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    return item


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))


def data_table(kind):
    values = COPY['tables'][kind]
    table = document.add_table(rows=0, cols=len(values[0]))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = {'results': [1.9, 1.25, 1.7, 1.95], 'quantities': [1.5, 3.1, 2.2], 'rubric': [5.2, 1.6]}[kind]
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for index, values_row in enumerate(values):
        row = table.add_row()
        prevent_row_split(row)
        for cell, value, width in zip(row.cells, values_row, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            item = cell.paragraphs[0]
            item.paragraph_format.space_before = Pt(4)
            item.paragraph_format.space_after = Pt(4)
            run = item.add_run(value)
            run.bold = index == 0
            run.font.size = Pt(10)
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    paragraph('', size=1, after=2)


def photographs(ids):
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column in table.columns:
        column.width = Inches(3.4)
    for start in range(0, len(ids), 2):
        batch = ids[start:start + 2]
        row = table.add_row()
        prevent_row_split(row)
        cells = [row.cells[0].merge(row.cells[1])] if len(batch) == 1 else row.cells
        for cell, photo_id in zip(cells, batch):
            PHOTO_ORDER.append(photo_id)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            item = cell.paragraphs[0]
            item.alignment = WD_ALIGN_PARAGRAPH.CENTER
            item.paragraph_format.space_before = Pt(6)
            item.paragraph_format.space_after = Pt(4)
            item.paragraph_format.keep_with_next = True
            picture = item.add_run().add_picture(str(ROOT / f'evidence/photo-{photo_id:02d}.jpg'), height=Inches(3.0))
            caption = COPY['captions'][str(photo_id)]
            picture._inline.docPr.set('descr', caption)
            picture._inline.docPr.set('title', f'Photo {photo_id:02d}')
            item = cell.add_paragraph(f'Foto {len(PHOTO_ORDER)}. {caption}')
            item.alignment = WD_ALIGN_PARAGRAPH.CENTER
            item.paragraph_format.space_after = Pt(10)
            item.paragraph_format.keep_together = True
            for run in item.runs:
                run.font.size = Pt(9.5)


paragraph(COPY['school'], size=12, bold=True, center=True, before=70, after=6)
paragraph(COPY['place'], size=11, center=True, after=6)
paragraph(COPY['subject'], size=11, center=True, after=72)
paragraph(COPY['title'], size=20, bold=True, center=True, after=8)
paragraph(COPY['subtitle'], size=17, bold=True, center=True, after=8)
paragraph('Construcción, pruebas y resultados', size=12, center=True, after=90)
paragraph(COPY['author'], size=13, bold=True, center=True, after=8)
paragraph(COPY['grade'], size=11, center=True, after=72)
paragraph(COPY['year'], size=11, center=True)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

for blocks in COPY['pages']:
    document.add_page_break()
    for kind, text in blocks:
        if kind == 'heading':
            item = paragraph(text, size=12, bold=True, before=9, after=7, keep=True)
            item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind == 'subheading':
            item = paragraph(text, bold=True, before=6, after=5, keep=True)
            item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif kind in COPY['tables']:
            data_table(kind)
        elif kind == 'photos':
            photographs(text)
        elif kind == 'equation':
            paragraph(text, center=True, after=9)
        elif kind == 'source':
            item = paragraph(text, size=10.5, after=6)
            item.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            paragraph(text)

assert sorted(PHOTO_ORDER) == list(range(1, 16))
document.save(TARGET)
(ROOT / 'reports').mkdir(exist_ok=True)
(ROOT / 'reports/photo-order.json').write_text(json.dumps(PHOTO_ORDER), encoding='utf-8')
print(f'Built {TARGET.name} with {len(PHOTO_ORDER)} original photographs.')
