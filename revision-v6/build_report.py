"""Create matching, plain school-report DOCX and PDF from one content source."""
import html
import json
from pathlib import Path

import fitz
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak

ROOT = Path(__file__).resolve().parent
COPY = json.loads((ROOT / 'document-copy.json').read_text(encoding='utf-8'))
PLAN = json.loads((ROOT / 'narration/revision-plan.json').read_text())
TARGET = ROOT.parent / 'Rocket-Results-and-Analysis-v6'
TABLES = {
    'results': [['Version', 'Height', 'Goal reached', 'Still needed']] + [[r['label'], f"About {r['heightMetres']:g} m", f"{r['goalPercent']}%", f"{r['shortfallMetres']:g} m"] for r in PLAN['results']],
    'quantities': [['Item', 'Calculation', 'Estimate'], ['Vinegar', '400 mL × ¼', '100 mL'], ['Baking soda', '50 g × ½', '25 g'], ['Dry rocket', 'Working assumption', '60 g']]
}


def main():
    for suffix in ('.docx', '.pdf'):
        assert not TARGET.with_suffix(suffix).exists(), TARGET.with_suffix(suffix)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size, normal.font.color.rgb = 'Arial', Pt(11), RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    document.core_properties.author = COPY['author']
    document.core_properties.title = COPY['title']
    document.core_properties.subject = COPY['subtitle']
    document.core_properties.language = 'en'
    pdfmetrics.registerFont(TTFont('SchoolArial', 'C:/Windows/Fonts/arial.ttf'))
    pdfmetrics.registerFont(TTFont('SchoolArialBold', 'C:/Windows/Fonts/arialbd.ttf'))
    pdfmetrics.registerFontFamily('SchoolArial', normal='SchoolArial', bold='SchoolArialBold')
    body = ParagraphStyle('Body', fontName='SchoolArial', fontSize=11, leading=12.65, spaceAfter=7, alignment=TA_JUSTIFY)
    heading = ParagraphStyle('Heading', parent=body, fontName='SchoolArialBold', fontSize=11, spaceBefore=10, spaceAfter=6, keepWithNext=True, alignment=0)
    title = ParagraphStyle('Title', parent=body, fontName='SchoolArialBold', fontSize=12, alignment=TA_CENTER, spaceAfter=5)
    author = ParagraphStyle('Author', parent=body, alignment=TA_CENTER, spaceAfter=13)
    cell_style = ParagraphStyle('Cell', parent=body, fontSize=10, leading=12, spaceAfter=0, alignment=0)
    story = []
    all_text = []

    def paragraph(text, kind='paragraph'):
        p = document.add_paragraph()
        run = p.add_run(text)
        if kind in ('heading', 'title'):
            run.bold = True
            p.paragraph_format.keep_with_next = True
        if kind in ('title', 'author'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == 'paragraph':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if kind == 'title':
            run.font.size = Pt(12)
        if kind == 'heading':
            p.paragraph_format.space_before = Pt(10)
        style = {'heading': heading, 'title': title, 'author': author}.get(kind, body)
        story.append(Paragraph(html.escape(text), style))
        all_text.append(text)

    paragraph(COPY['title'], 'title')
    paragraph(COPY['subtitle'], 'title')
    paragraph(COPY['author'], 'author')
    for page_index, blocks in enumerate(COPY['pages']):
        if page_index:
            document.add_page_break()
            story.append(PageBreak())
        for kind, text in blocks:
            if kind not in TABLES:
                paragraph(text, kind)
                continue
            rows = TABLES[kind]
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for run in p.runs:
                            run.font.size = Pt(10)
                            run.bold = index == 0
                if index == 0:
                    repeat = OxmlElement('w:tblHeader')
                    repeat.set(qn('w:val'), 'true')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            widths = [155, 105, 115, 114.6] if kind == 'results' else [150, 205, 134.6]
            data = [[Paragraph(('<b>' if i == 0 else '') + html.escape(cell) + ('</b>' if i == 0 else ''), cell_style) for cell in row] for i, row in enumerate(rows)]
            pdf_table = Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
            pdf_table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .5, colors.black), ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'), ('TOPPADDING', (0, 0), (-1, -1), 6), ('BOTTOMPADDING', (0, 0), (-1, -1), 6)]))
            story.extend([pdf_table, Spacer(1, 9)])
            all_text.extend(cell for row in rows for cell in row)
    document.save(TARGET.with_suffix('.docx'))
    SimpleDocTemplate(str(TARGET.with_suffix('.pdf')), pagesize=letter, topMargin=61.2, bottomMargin=61.2,
                      leftMargin=61.2, rightMargin=61.2, title=COPY['title'], author=COPY['author']).build(story)
    pdf = fitz.open(TARGET.with_suffix('.pdf'))
    assert len(pdf) == 2, len(pdf)
    extracted = ' '.join(' '.join(page.get_text().split()) for page in pdf)
    for text in all_text:
        assert ' '.join(text.split()) in extracted, text
    for index, page in enumerate(pdf):
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(ROOT / f'out/review/report-page-{index + 1}.png')
    check = Document(TARGET.with_suffix('.docx'))
    assert len(check.tables) == 2
    report = {'pages': len(pdf), 'wordCount': len(' '.join(all_text).split()), 'allTextPresentInPdf': True,
              'docxParsed': True, 'style': 'Black Arial, centered 12-point title, 11-point body, plain grid tables, no decorative backgrounds or footer',
              'referencesInspected': ['Descartes final activity', 'Hydrocarbons activity', 'Tribunal report generator'],
              'language': 'English per latest session protocol', 'files': [str(TARGET.with_suffix('.docx')), str(TARGET.with_suffix('.pdf'))]}
    (ROOT / 'reports/document-build.json').write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))


if __name__ == '__main__':
    main()
